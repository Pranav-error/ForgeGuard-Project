# Build ForgeGuard_IEEE_Paper.docx from paper_out.html (run paper_build.py first).
# Figures are rendered to PNG with headless Chrome, text goes through pandoc, layout is fixed with python-docx.
import re, subprocess, pathlib, copy
from PIL import Image, ImageChops
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
h = pathlib.Path("paper_out.html").read_text()
style = re.search(r"<style>.*?</style>", h, re.S).group(0)
fonts = re.search(r'<link href="https://fonts[^>]*>', h).group(0)

# 1. figures -> PNG (render each diagram at column width, 3x scale, trim whitespace)
figs = re.findall(r'<div class="dg">.*?</div>\s*</div>\s*<div class="cap">', h, re.S)
for i, f in enumerate(figs, 1):
    dg = f[: f.rfind('<div class="cap">')]
    pathlib.Path(f"fig{i}.html").write_text(f"<html><head>{fonts}{style}</head><body style='width:330px;padding:4px'>{dg}</body></html>")
    subprocess.run([CHROME, "--headless", "--disable-gpu", "--hide-scrollbars", "--force-device-scale-factor=3",
                    "--window-size=340,600", "--virtual-time-budget=5000", f"--screenshot=fig{i}.png", f"fig{i}.html"],
                   check=True, capture_output=True)
    im = Image.open(f"fig{i}.png").convert("RGB")
    im.crop(ImageChops.difference(im, Image.new("RGB", im.size, "white")).getbbox()).save(f"fig{i}.png")
    pathlib.Path(f"fig{i}.html").unlink()
for i, f in enumerate(figs, 1):
    h = h.replace(f[: f.rfind('<div class="cap">')], f'<p class="FIG{i}">FIGURE{i}</p>', 1)

# 2. simplify markup for pandoc: markers for title/authors/abstract/captions
h = re.sub(r'<div class="title">(.*?)</div>', r'<p>@@TITLE@@\1</p>', h, flags=re.S)
h = re.sub(r'<div class="n">(.*?)</div>', r'<p>@@AUTH@@\1</p><p>@@AFF@@', h)
h = re.sub(r'<div class="cap">(.*?)</div>', r'<p>@@CAP@@\1</p>', h, flags=re.S)
h = re.sub(r'<div class="tcap">(.*?)<br>(.*?)</div>', r'<p>@@TCAP@@\1</p><p>@@TCAP@@\2</p>', h, flags=re.S)
h = h.replace('<p class="abs">', '<p>@@ABS@@').replace('<p class="kw">', '<p>@@KW@@')
h = h.replace('<div class="tbl wide">', '<p>@@WIDE_START@@</p><div>')
h = re.sub(r'(@@WIDE_START@@.*?</table>\s*</div>)', r'\1<p>@@WIDE_END@@</p>', h, flags=re.S)
h = re.sub(r'<p><span>(\[\d+\])</span>', r'<p>@@REF@@\1 ', h)
h = h.replace("<br>REVA", " REVA")  # affiliation on one line, split later
pathlib.Path("_pandoc.html").write_text(h)
subprocess.run(["pandoc", "_pandoc.html", "-o", "_raw.docx"], check=True)
pathlib.Path("_pandoc.html").unlink()

# 3. layout
d = Document("_raw.docx")
pathlib.Path("_raw.docx").unlink()
for s in d.styles:
    try:
        s.font.name = "Times New Roman"; s.font.color.rgb = None
        s.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    except Exception: pass

def fmt(p, size=10, bold=None, italic=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0, indent=None):
    pf = p.paragraph_format
    pf.alignment, pf.space_before, pf.space_after, pf.line_spacing = align, Pt(before), Pt(after), 1.0
    if indent is not None: pf.first_line_indent = Pt(indent)
    for r in p.runs:
        r.font.size, r.font.name = Pt(size), "Times New Roman"
        if bold is not None: r.font.bold = bold
        if italic is not None: r.font.italic = italic

def strip(p, tag):
    for r in p.runs:
        if tag in r.text: r.text = r.text.replace(tag, ""); break

def cols(sectPr, n):
    c = sectPr.find(qn("w:cols"))
    if c is None: c = OxmlElement("w:cols"); sectPr.append(c)
    c.set(qn("w:num"), str(n)); c.set(qn("w:space"), "346")

def section_break_after(p, ncols):
    """End a section at paragraph p; that section gets ncols columns."""
    pPr = p._p.get_or_add_pPr()
    sp = copy.deepcopy(d.sections[-1]._sectPr)
    t = sp.find(qn("w:type"))
    if t is None: t = OxmlElement("w:type"); sp.insert(0, t)
    t.set(qn("w:val"), "continuous")
    cols(sp, ncols); pPr.append(sp)

sec = d.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.62)
sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(0.9)

paras = list(d.paragraphs)
last_head = None
for p in paras:
    t = p.text
    sn = p.style.name
    if "@@TITLE@@" in t:
        strip(p, "@@TITLE@@"); fmt(p, 22, False, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    elif "@@AUTH@@" in t:
        strip(p, "@@AUTH@@"); fmt(p, 10.5, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=1)
    elif "@@AFF@@" in t:
        strip(p, "@@AFF@@"); fmt(p, 9, align=WD_ALIGN_PARAGRAPH.CENTER, after=2); last_head = p
    elif "@@ABS@@" in t or "@@KW@@" in t:
        tag = "@@ABS@@" if "@@ABS@@" in t else "@@KW@@"
        strip(p, tag); fmt(p, 9, True, indent=10, after=4)
        if last_head is not None: section_break_after(last_head, 1); last_head = None
    elif "@@CAP@@" in t:
        strip(p, "@@CAP@@"); fmt(p, 8, align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=8)
    elif "@@TCAP@@" in t:
        strip(p, "@@TCAP@@")
        for r in p.runs: r.font.small_caps = True
        fmt(p, 8, align=WD_ALIGN_PARAGRAPH.CENTER, before=6)
    elif "@@REF@@" in t:
        strip(p, "@@REF@@"); fmt(p, 8, align=WD_ALIGN_PARAGRAPH.LEFT, after=2, indent=-16)
        p.paragraph_format.left_indent = Pt(16)
    elif re.fullmatch(r"FIGURE\d", t):
        i = t[-1]; p.text = ""
        p.add_run().add_picture(f"fig{i}.png", width=Inches(3.3))
        fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=6)
    elif "@@WIDE_START@@" in t:
        prev = p._p.getprevious()
        p.text = ""; section_break_after(p, 2)
    elif "@@WIDE_END@@" in t:
        p.text = ""; section_break_after(p, 1)
    elif sn.startswith("Heading 2") or sn == "Heading 2":
        fmt(p, 12, True, align=WD_ALIGN_PARAGRAPH.LEFT, before=9, after=4)
    elif sn.startswith("Heading 3"):
        fmt(p, 10, False, True, align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=2)
    elif sn.startswith("Heading 1"):
        fmt(p, 22, align=WD_ALIGN_PARAGRAPH.CENTER)
    elif "List" in sn:
        fmt(p, 10, after=2)
    else:
        fmt(p, 10, indent=10)
# body after last section break is two-column and continues on the same page
fs = d.sections[-1]._sectPr; cols(fs, 2)
t = fs.find(qn("w:type"))
if t is None: t = OxmlElement("w:type"); fs.insert(0, t)
t.set(qn("w:val"), "continuous")

# pandoc turns <title> into a Title paragraph: drop it
for p in d.paragraphs:
    if p.style.name in ("Title", "Subtitle"): p._p.getparent().remove(p._p)

# authors: 2-column borderless grid instead of a stacked list
auth = [p for p in d.paragraphs if p.text.startswith("[Author") or p.text.startswith("School of Computer")]
if auth:
    pairs = [(auth[i].text, auth[i + 1].text) for i in range(0, len(auth), 2)]
    g = d.add_table(rows=(len(pairs) + 1) // 2, cols=2)
    for k, (n, a) in enumerate(pairs):
        c = g.cell(k // 2, k % 2); c.text = ""
        pn = c.paragraphs[0]; pn.add_run(n); fmt(pn, 10.5, align=WD_ALIGN_PARAGRAPH.CENTER, before=4)
        for line in a.replace(" REVA", "\nREVA").split("\n"):
            pa = c.add_paragraph(line); fmt(pa, 9, align=WD_ALIGN_PARAGRAPH.CENTER)
    last = auth[-1]._p
    last.addprevious(g._tbl)
    for p in auth[:-1]: p._p.getparent().remove(p._p)
    for r in last.findall(qn("w:r")): last.remove(r)   # keep the paragraph: it carries the section break
    g.alignment = 1

for p in d.paragraphs:
    if "List" in p.style.name:
        p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Pt(12), Pt(-9)
    if p.text.startswith("Table ") and len(p.text) < 9 or (p.runs and p.runs[0].font.small_caps):
        p.paragraph_format.keep_with_next = True

# tables: booktabs-like rules, small font
for tb in d.tables[1:]:  # first table is the author grid
    tb.autofit = True
    for ri, row in enumerate(tb.rows):
        trPr = row._tr.get_or_add_trPr(); trPr.append(OxmlElement("w:cantSplit"))
        for c in row.cells:
            for p in c.paragraphs:
                fmt(p, 7.5, align=WD_ALIGN_PARAGRAPH.LEFT)
                p.paragraph_format.keep_with_next = ri < len(tb.rows) - 1
    tblPr = tb._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side, sz in (("top", 8), ("bottom", 8), ("insideH", 0)):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "single" if sz else "nil"); e.set(qn("w:sz"), str(sz)); b.append(e)
    for old in tblPr.findall(qn("w:tblBorders")): tblPr.remove(old)
    tblPr.append(b)
    for c in tb.rows[0].cells:  # rule under header
        tcPr = c._tc.get_or_add_tcPr(); tb_ = OxmlElement("w:tcBorders")
        e = OxmlElement("w:bottom"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); tb_.append(e); tcPr.append(tb_)

# don't balance columns before the full-width table (keeps Table 2 in one piece; Table 3 starts a page)
st = d.settings.element
cp = st.find(qn("w:compat"))
if cp is None: cp = OxmlElement("w:compat"); st.append(cp)
cp.insert(0, OxmlElement("w:noColumnBalance"))
# pandoc bullets: tighter hanging indent
for p in d.paragraphs:
    if p._p.pPr is not None and p._p.pPr.numPr is not None:
        p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Pt(12), Pt(-9)
d.save("ForgeGuard_IEEE_Paper.docx")
for i in range(1, len(figs) + 1): pathlib.Path(f"fig{i}.png").unlink()
print("ok")
