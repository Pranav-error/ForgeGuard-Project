# Build System_Design_Document.docx from design_out.html (run design_build.py first).
# The inline SVG diagram can't survive a plain pandoc conversion, so it's rendered to
# PNG with headless Chrome first and swapped in as an <img>.
import re, subprocess, pathlib
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
h = pathlib.Path("design_out.html").read_text()

svg = re.search(r'<svg class="diagram".*?</svg>', h, re.S).group(0)
pathlib.Path("_diagram.html").write_text(f"<html><body style='margin:0;padding:8px;background:#fff'>{svg}</body></html>")
subprocess.run([CHROME, "--headless", "--disable-gpu", "--hide-scrollbars", "--force-device-scale-factor=3",
                "--window-size=800,700", "--virtual-time-budget=5000", "--screenshot=_diagram.png", "_diagram.html"],
               check=True, capture_output=True)
pathlib.Path("_diagram.html").unlink()
h = h.replace(svg, '<img src="_diagram.png">')
pathlib.Path("_design_pandoc.html").write_text(h)

subprocess.run(["pandoc", "_design_pandoc.html", "-o", "_design_raw.docx"], check=True)
pathlib.Path("_design_pandoc.html").unlink()

d = Document("_design_raw.docx")
pathlib.Path("_design_raw.docx").unlink()

for s in d.styles:
    try:
        s.font.name = "Calibri"
        s.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    except Exception:
        pass

# pandoc's leading Title paragraph duplicates the in-body H1: drop it
for p in list(d.paragraphs):
    if p.style.name == "Title":
        p._p.getparent().remove(p._p)

for p in d.paragraphs:
    sn = p.style.name
    if sn.startswith("Heading 1"):
        for r in p.runs: r.font.size, r.font.bold = Pt(16), True
    elif sn.startswith("Heading 2"):
        for r in p.runs: r.font.size, r.font.bold = Pt(13), True
        p.paragraph_format.space_before = Pt(10)
    elif sn.startswith("Heading 3"):
        for r in p.runs: r.font.size, r.font.bold = Pt(11.5), True
        p.paragraph_format.space_before = Pt(8)
    else:
        for r in p.runs: r.font.size = Pt(11)

for t in d.tables:
    t.autofit = True
    tblPr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "000000")
        borders.append(e)
    tblPr.append(borders)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs: r.font.size = Pt(10)

for p in d.paragraphs:
    for r in p.runs:
        if r._element.find(qn("w:drawing")) is not None:
            continue
    if p.runs and p.runs[0]._element.find(qn("w:drawing")) is not None:
        # resize the inline image to the text width and center it
        inline = p.runs[0]._element.findall(".//" + qn("wp:inline"))
        for shape in inline:
            ext = shape.find(qn("wp:extent"))
            w, hgt = int(ext.get("cx")), int(ext.get("cy"))
            new_w = Inches(6.3)
            scale = new_w / w
            ext.set("cx", str(new_w)); ext.set("cy", str(int(hgt * scale)))
        p.alignment = 1

d.save("System_Design_Document.docx")
pathlib.Path("_diagram.png").unlink()
print("ok")
