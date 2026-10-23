# Build Project_Objectives.docx from objectives_out.html (run objectives_build.py first).
import subprocess, pathlib
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

subprocess.run(["pandoc", "objectives_out.html", "-o", "_obj_raw.docx"], check=True)

d = Document("_obj_raw.docx")
pathlib.Path("_obj_raw.docx").unlink()

for s in d.styles:
    try:
        s.font.name = "Calibri"
        s.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    except Exception:
        pass

# pandoc turns <title> into a redundant leading Title paragraph: drop it
for p in list(d.paragraphs):
    if p.style.name == "Title":
        p._p.getparent().remove(p._p)

for p in d.paragraphs:
    sn = p.style.name
    if sn.startswith("Heading 1"):
        for r in p.runs: r.font.size, r.font.bold = Pt(16), True
    elif sn.startswith("Heading 2"):
        for r in p.runs: r.font.size, r.font.bold = Pt(13), True
        p.paragraph_format.space_before = Pt(10)
    else:
        for r in p.runs: r.font.size = Pt(11)

from docx.oxml import OxmlElement
for t in d.tables:
    t.autofit = True
    tblPr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "000000")
        borders.append(e)
    tblPr.append(borders)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs: r.font.size = Pt(10)

d.save("Project_Objectives.docx")
print("ok")
